import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os, shutil

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_resume_docx():
    doc = docx.Document()

    # Set page margins to 0.75 in
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    PRIMARY_COLOR = RGBColor(2, 132, 199)    # #0284c7 (Ocean Blue)
    SECONDARY_COLOR = RGBColor(15, 23, 42)   # #0f172a (Dark Slate)
    MUTED_COLOR = RGBColor(100, 116, 139)    # #64748b (Slate Muted)

    # 1. Header Section (Table with Photo & Contact)
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False

    # Widths
    header_table.columns[0].width = Inches(1.5)
    header_table.columns[1].width = Inches(5.5)

    cell_photo = header_table.cell(0, 0)
    cell_info = header_table.cell(0, 1)

    # Insert Profile Photo
    photo_path = r'd:\01_Projects\antigravity\sample\profile.jpg'
    if os.path.exists(photo_path):
        p_photo = cell_photo.paragraphs[0]
        p_photo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_img = p_photo.add_run()
        run_img.add_picture(photo_path, width=Inches(1.35))

    # Info Column
    p_name = cell_info.paragraphs[0]
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(4)
    r_name = p_name.add_run("ABDUL RASAK M R")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(22)
    r_name.font.bold = True
    r_name.font.color.rgb = SECONDARY_COLOR

    p_title = cell_info.add_paragraph()
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("Senior Software Engineer | Healthcare Integration & Microservices Specialist")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(12)
    r_title.font.bold = True
    r_title.font.color.rgb = PRIMARY_COLOR

    p_contact = cell_info.add_paragraph()
    p_contact.paragraph_format.space_after = Pt(0)
    r_contact = p_contact.add_run("📱 Mobile: +91 9447723429  |  📧 Email: rasakmar@gmail.com\n📍 Location: Muvattupuzha, Kerala, India  |  🎓 B.Tech CSE (MG University 2012)  |  💼 11+ Yrs Exp")
    r_contact.font.name = "Calibri"
    r_contact.font.size = Pt(10)
    r_contact.font.color.rgb = MUTED_COLOR

    # Add Divider Line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(8)
    p_div.paragraph_format.space_after = Pt(12)
    r_div = p_div.add_run("―" * 58)
    r_div.font.color.rgb = PRIMARY_COLOR
    r_div.font.bold = True

    # Helper function for headings
    def add_section_heading(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title_text)
        r.font.name = "Calibri"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = PRIMARY_COLOR
        return p

    # 2. Executive Summary
    add_section_heading("PROFESSIONAL SUMMARY")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_after = Pt(10)
    p_sum.paragraph_format.line_spacing = 1.15
    r_sum = p_sum.add_run(
        "Results-driven Senior Software Engineer with over 11+ years of hands-on experience architecting and engineering "
        "enterprise-grade Healthcare Information Systems (HIS/CMS), Pharmacy Systems, and Supply Chain ERPs across UAE and India. "
        "Recognized domain specialist in regional healthcare interoperability (HAAD, DHA, Malafi, Nabidh, Riayati) leveraging "
        "HL7 v2, FHIR messaging, and Mirth Connect. Proven track record in Laboratory Information System (LIS) device integrations. "
        "Currently leading cloud-native microservices modernization using Spring Boot, React.js, OAuth2, and PostgreSQL."
    )
    r_sum.font.name = "Calibri"
    r_sum.font.size = Pt(10.5)
    r_sum.font.color.rgb = SECONDARY_COLOR

    # 3. Technical Skills Table
    add_section_heading("CORE COMPETENCIES & TECHNICAL SKILLS")
    skills_data = [
        ("Backend & Microservices", "Java (7/8/17+), Spring Boot, Spring MVC, Spring Security, Microservices Architecture, OAuth2, RESTful APIs"),
        ("Healthcare Interoperability", "HL7 v2 Standard, FHIR Messaging, Mirth Connect, HAAD (Abu Dhabi), DHA (Dubai), Malafi, Nabidh, Riayati, LIS Machine Integration"),
        ("Frontend & Web", "React.js, JavaScript (ES6+), jQuery, HTML5, CSS3, Responsive Web Design"),
        ("Databases & Reporting", "MySQL (5.6+), PostgreSQL, Jasper Reports, SQL Optimization"),
        ("Domain Expertise", "Hospital Information Systems (HIS), Clinical Management (CMS), Pharmacy Systems, Revenue Cycle Management (RCM), iERP (Procurement & Warehouse)")
    ]

    skills_table = doc.add_table(rows=len(skills_data)+1, cols=2)
    skills_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    skills_table.autofit = False
    skills_table.columns[0].width = Inches(2.2)
    skills_table.columns[1].width = Inches(4.8)

    # Header Row
    hdr_cells = skills_table.rows[0].cells
    set_cell_background(hdr_cells[0], "0284C7")
    set_cell_background(hdr_cells[1], "0284C7")

    r_h0 = hdr_cells[0].paragraphs[0].add_run("Category")
    r_h0.font.bold = True
    r_h0.font.color.rgb = RGBColor(255, 255, 255)
    r_h0.font.name = "Calibri"

    r_h1 = hdr_cells[1].paragraphs[0].add_run("Technologies & Frameworks")
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(255, 255, 255)
    r_h1.font.name = "Calibri"

    for idx, (cat, tech) in enumerate(skills_data):
        row_cells = skills_table.rows[idx+1].cells
        if idx % 2 == 1:
            set_cell_background(row_cells[0], "F8FAFC")
            set_cell_background(row_cells[1], "F8FAFC")
        
        r0 = row_cells[0].paragraphs[0].add_run(cat)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = SECONDARY_COLOR

        r1 = row_cells[1].paragraphs[0].add_run(tech)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = SECONDARY_COLOR

    # 4. Work Experience
    add_section_heading("WORK EXPERIENCE")

    # Role 1
    p_r1 = doc.add_paragraph()
    p_r1.paragraph_format.space_before = Pt(8)
    p_r1.paragraph_format.space_after = Pt(2)
    r_r1_title = p_r1.add_run("Senior Software Engineer")
    r_r1_title.font.name = "Calibri"
    r_r1_title.font.size = Pt(12)
    r_r1_title.font.bold = True
    r_r1_title.font.color.rgb = SECONDARY_COLOR

    r_r1_comp = p_r1.add_run("  |  Safecare Technologies Pvt Ltd, Muvattupuzha")
    r_r1_comp.font.name = "Calibri"
    r_r1_comp.font.size = Pt(11)
    r_r1_comp.font.color.rgb = PRIMARY_COLOR

    p_r1_date = doc.add_paragraph()
    p_r1_date.paragraph_format.space_after = Pt(4)
    r_r1_d = p_r1_date.add_run("March 2015 – Present (11+ Years)   [Career Progression: Junior Software Engineer (2.1 yrs) ➔ Software Engineer (1 yr) ➔ Senior Software Engineer (2018 – Present)]")
    r_r1_d.font.name = "Calibri"
    r_r1_d.font.size = Pt(9.5)
    r_r1_d.font.italic = True
    r_r1_d.font.color.rgb = MUTED_COLOR

    bullets = [
        "Architecture Modernization & Microservices: Spearheading the migration of monolithic Java/Spring MVC suites into modern Spring Boot Microservices, React.js frontends, OAuth2 authentication, and PostgreSQL databases.",
        "Healthcare Interoperability (HIE Integrations): Engineered bi-directional integration engines connecting Cortex HIS/CMS with UAE regional Health Information Exchanges including Malafi (Abu Dhabi), Nabidh (Dubai), and Riayati (MOHAP) using HL7 v2, FHIR standards, and Mirth Connect.",
        "Laboratory Machine Integration (LIS): Implemented automated laboratory machine connectivity interfacing medical auto-analyzers directly with LIS via HL7 protocols.",
        "Healthcare Claims & RCM Automation: Designed end-to-end e-Claim processing workflows for HAAD & DHA regulatory compliance (Prior Approvals, Initial Submissions, Remittance Advice parsing, and Resubmissions).",
        "iERP & Mobile API Development: Built custom procurement and warehouse management suite (iERP), developing REST APIs connecting mobile applications to track real-time field worker activities.",
        "Reporting & Analytics: Engineered clinical, inventory, and financial reporting modules utilizing Jasper Reports."
    ]

    for b in bullets:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_after = Pt(3)
        p_b.paragraph_format.line_spacing = 1.15
        r_b = p_b.add_run(b)
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(10)
        r_b.font.color.rgb = SECONDARY_COLOR

    # Role 2
    p_r2 = doc.add_paragraph()
    p_r2.paragraph_format.space_before = Pt(10)
    p_r2.paragraph_format.space_after = Pt(2)
    r_r2_t = p_r2.add_run("Freelance Web Designer & Developer")
    r_r2_t.font.name = "Calibri"
    r_r2_t.font.size = Pt(11.5)
    r_r2_t.font.bold = True
    r_r2_t.font.color.rgb = SECONDARY_COLOR
    p_r2.add_run("   (Jan 2013 – Feb 2015)").font.color.rgb = MUTED_COLOR

    p_b2 = doc.add_paragraph(style='List Bullet')
    p_b2.paragraph_format.space_after = Pt(3)
    r_b2 = p_b2.add_run("Designed and developed custom, responsive websites and client web applications using HTML5, CSS3, and JavaScript.")
    r_b2.font.name = "Calibri"
    r_b2.font.size = Pt(10)

    # Role 3
    p_r3 = doc.add_paragraph()
    p_r3.paragraph_format.space_before = Pt(8)
    p_r3.paragraph_format.space_after = Pt(2)
    r_r3_t = p_r3.add_run("Java Software Developer Intern")
    r_r3_t.font.name = "Calibri"
    r_r3_t.font.size = Pt(11.5)
    r_r3_t.font.bold = True
    r_r3_t.font.color.rgb = SECONDARY_COLOR
    p_r3.add_run("   | Brazier IT Solutions, Bangalore (Sep 2012 – Jan 2013)").font.color.rgb = MUTED_COLOR

    p_b3 = doc.add_paragraph(style='List Bullet')
    p_b3.paragraph_format.space_after = Pt(3)
    r_b3 = p_b3.add_run("Hands-on core Java application development, OOP principles, SQL database queries, and module testing.")
    r_b3.font.name = "Calibri"
    r_b3.font.size = Pt(10)

    # 5. Major Projects
    add_section_heading("MAJOR FEATURED PROJECTS")

    projects = [
        ("Cortex HIS & CMS", "Integrated UAE HIEs (Malafi, Nabidh, Riayati) and regional health authorities (HAAD, DHA) via HL7, FHIR, and Mirth Connect. (Java 8, Spring MVC, MySQL, Jasper Reports)."),
        ("Microservices Migration & LIS Platform", "Architectural refactoring to Spring Boot Microservices, React.js, OAuth2, and PostgreSQL. Automated medical lab machine integration via HL7."),
        ("Pharmacy Management System (UAE & India)", "Comprehensive drug dispensing and inventory system with HAAD & MOH claim prior approvals, submissions, remittance collection, and resubmissions."),
        ("iERP (Procurement, Supply Chain & Warehouse)", "Procurement and warehouse management platform with RESTful APIs supporting mobile app task tracking for field workers.")
    ]

    for proj_title, proj_desc in projects:
        p_p = doc.add_paragraph()
        p_p.paragraph_format.space_before = Pt(4)
        p_p.paragraph_format.space_after = Pt(2)
        r_pt = p_p.add_run(f"•  {proj_title}: ")
        r_pt.font.name = "Calibri"
        r_pt.font.size = Pt(10)
        r_pt.font.bold = True
        r_pt.font.color.rgb = PRIMARY_COLOR

        r_pd = p_p.add_run(proj_desc)
        r_pd.font.name = "Calibri"
        r_pd.font.size = Pt(10)
        r_pd.font.color.rgb = SECONDARY_COLOR

    # 6. Education
    add_section_heading("EDUCATION")
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_after = Pt(4)
    r_edu_title = p_edu.add_run("Bachelor of Technology (B.Tech) in Computer Science & Engineering")
    r_edu_title.font.name = "Calibri"
    r_edu_title.font.size = Pt(11)
    r_edu_title.font.bold = True
    r_edu_title.font.color.rgb = SECONDARY_COLOR

    p_edu_sub = doc.add_paragraph()
    r_edu_sub = p_edu_sub.add_run("MG University, Kerala, India  |  Passout Year: 2012")
    r_edu_sub.font.name = "Calibri"
    r_edu_sub.font.size = Pt(10)
    r_edu_sub.font.color.rgb = MUTED_COLOR

    out_path = r'd:\01_Projects\antigravity\sample\Abdul_Rasak_M_R_Resume.docx'
    doc.save(out_path)
    print("Saved Word Resume successfully to:", out_path)

    # Copy to artifacts directory
    artifact_dir = r'C:\Users\admin\.gemini\antigravity\brain\7844875c-6066-4c3a-a660-4c854f455799'
    shutil.copy(out_path, os.path.join(artifact_dir, 'Abdul_Rasak_M_R_Resume.docx'))

if __name__ == '__main__':
    create_resume_docx()
